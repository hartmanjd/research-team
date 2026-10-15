"""Convert a DOL disclosure workbook's first sheet to CSV, streaming row by row.

The disclosure files are large (up to ~150 MB each), so the sheet is read in
openpyxl's read-only mode instead of being loaded whole. Fully blank rows,
which DOL pads the sheets with, are skipped. Writes OUT.csv.manifest.json with
the workbook as input.

Usage: python3 xlsx_to_csv.py IN.xlsx OUT.csv [--source-url URL]
"""
import argparse, csv, datetime, os, sys
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from provenance import write_manifest


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time() else v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def convert(xlsx, out):
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    rows = wb.worksheets[0].iter_rows(values_only=True)
    n = 0
    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        header = [cell_text(v).strip() for v in next(rows)]
        w.writerow(header)
        for r in rows:
            vals = [cell_text(v) for v in r][:len(header)]
            if any(vals):
                w.writerow(vals + [""] * (len(header) - len(vals)))
                n += 1
    wb.close()
    return n


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("xlsx")
    ap.add_argument("out")
    ap.add_argument("--source-url", action="append", default=[])
    a = ap.parse_args(argv)
    n = convert(a.xlsx, a.out)
    write_manifest(a.out, __file__, inputs=[a.xlsx], sources=[{"url": u} for u in a.source_url])
    print(f"{n} rows -> {a.out}")


if __name__ == "__main__":
    main()
